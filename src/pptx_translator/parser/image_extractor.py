import hashlib
import io
import logging
from typing import List

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

logger = logging.getLogger(__name__)


class ExtractedImageInfo:
    def __init__(
        self,
        image_id: str,
        slide_index: int,
        shape_id: int,
        image_bytes: bytes,
        ext: str,
        slide_bbox: List[float],  # [left, top, width, height] in EMU
        crop: List[float],  # [crop_left, crop_top, crop_right, crop_bottom] normalized (0.0 - 1.0)
        rotation: float,  # degrees
        image_hash: str,
        pixel_width: int,
        pixel_height: int,
    ):
        self.image_id = image_id
        self.slide_index = slide_index
        self.shape_id = shape_id
        self.image_bytes = image_bytes
        self.ext = ext
        self.slide_bbox = slide_bbox
        self.crop = crop
        self.rotation = rotation
        self.image_hash = image_hash
        self.pixel_width = pixel_width
        self.pixel_height = pixel_height


class ImageExtractor:
    """Extracts images from PPTX slides with crop, rotation, and slide bbox info."""

    def __init__(self, pptx_path: str):
        self.pptx_path = pptx_path

    def extract_images_from_slide(self, slide, slide_index: int) -> List[ExtractedImageInfo]:
        extracted: List[ExtractedImageInfo] = []
        for shape in slide.shapes:
            self._process_shape(shape, slide_index, extracted)
        return extracted

    def _process_shape(self, shape, slide_index: int, extracted: List[ExtractedImageInfo]) -> None:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub_shape in shape.shapes:
                self._process_shape(sub_shape, slide_index, extracted)
            return

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
            hasattr(shape, "image") and shape.image is not None
        ):
            try:
                image_bytes = shape.image.blob
                ext = shape.image.ext
                image_hash = hashlib.sha256(image_bytes).hexdigest()

                pil_img = Image.open(io.BytesIO(image_bytes))
                pw, ph = pil_img.size

                crop_left = float(getattr(shape, "crop_left", 0.0) or 0.0)
                crop_top = float(getattr(shape, "crop_top", 0.0) or 0.0)
                crop_right = float(getattr(shape, "crop_right", 0.0) or 0.0)
                crop_bottom = float(getattr(shape, "crop_bottom", 0.0) or 0.0)
                rotation = float(getattr(shape, "rotation", 0.0) or 0.0)

                shape_id = shape.shape_id if hasattr(shape, "shape_id") else 0
                image_id = f"slide_{slide_index+1}_image_{shape_id}"

                slide_bbox = [
                    float(shape.left) if hasattr(shape, "left") else 0.0,
                    float(shape.top) if hasattr(shape, "top") else 0.0,
                    float(shape.width) if hasattr(shape, "width") else 0.0,
                    float(shape.height) if hasattr(shape, "height") else 0.0,
                ]

                info = ExtractedImageInfo(
                    image_id=image_id,
                    slide_index=slide_index,
                    shape_id=shape_id,
                    image_bytes=image_bytes,
                    ext=ext,
                    slide_bbox=slide_bbox,
                    crop=[crop_left, crop_top, crop_right, crop_bottom],
                    rotation=rotation,
                    image_hash=image_hash,
                    pixel_width=pw,
                    pixel_height=ph,
                )
                extracted.append(info)
            except Exception as e:
                logger.warning(f"Failed to extract image from shape on slide {slide_index+1}: {e}")
