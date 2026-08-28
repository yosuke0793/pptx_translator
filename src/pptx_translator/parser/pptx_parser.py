import logging
from typing import List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_translator.cache.cache_manager import CacheManager
from pptx_translator.models.element import ElementStyle, TextElement
from pptx_translator.models.slide import IntermediateRepresentation, SlideModel
from pptx_translator.ocr.detector import OCRDetector, OCRResultItem
from pptx_translator.ocr.target_classifier import TargetClassifier
from pptx_translator.ocr.text_grouper import TextGrouper
from pptx_translator.parser.image_extractor import ImageExtractor
from pptx_translator.parser.text_run_merger import is_translatable_text, merge_paragraph_runs

logger = logging.getLogger(__name__)


class PPTXParser:
    def __init__(
        self,
        pptx_path: str,
        translate_notes: bool = True,
        ocr_lang: str = "japan",
        ocr_confidence_threshold: float = 0.90,
        cache_manager: Optional[CacheManager] = None,
    ):
        self.pptx_path = pptx_path
        self.translate_notes = translate_notes
        self.ocr_lang = ocr_lang
        self.ocr_confidence_threshold = ocr_confidence_threshold
        self.cache_manager = cache_manager or CacheManager()
        self.prs = Presentation(pptx_path)

        self.image_extractor = ImageExtractor(pptx_path)
        self.ocr_detector = OCRDetector(lang=ocr_lang)
        self.text_grouper = TextGrouper()
        self.target_classifier = TargetClassifier()

    def parse(self) -> IntermediateRepresentation:
        slides: List[SlideModel] = []
        presentation_title: Optional[str] = None

        for slide_idx, slide in enumerate(self.prs.slides):
            slide_elements: List[TextElement] = []
            slide_title: Optional[str] = None

            # Check for slide title
            if slide.shapes.title and slide.shapes.title.text_frame:
                slide_title = slide.shapes.title.text_frame.text.strip()
                if slide_idx == 0 and not presentation_title:
                    presentation_title = slide_title

            # Parse shapes
            for shape in slide.shapes:
                self._parse_shape(shape, slide_idx, slide_elements)

            # Extract images & OCR
            extracted_images = self.image_extractor.extract_images_from_slide(slide, slide_idx)
            for img_info in extracted_images:
                ocr_items = self._get_ocr_for_image(img_info)
                grouped_items = self.text_grouper.group(ocr_items)

                for item_idx, item in enumerate(grouped_items):
                    is_target, reason = self.target_classifier.is_target(
                        item, min_confidence=self.ocr_confidence_threshold
                    )
                    elem_id = f"{img_info.image_id}_text_{item_idx+1}"

                    X, Y, W, H = img_info.slide_bbox
                    Wi, Hi = max(1, img_info.pixel_width), max(1, img_info.pixel_height)
                    xi, yi, wi, hi = item.bbox

                    xp = X + (xi / Wi) * W
                    yp = Y + (yi / Hi) * H
                    wp = (wi / Wi) * W
                    hp = (hi / Hi) * H
                    slide_bbox = [xp, yp, wp, hp]

                    style = ElementStyle(
                        color=item.text_color,
                        background_color=item.background_color,
                    )

                    element = TextElement(
                        element_id=elem_id,
                        slide_index=slide_idx,
                        source_type="image_text",
                        source_text=item.text,
                        bbox=slide_bbox,
                        parent_image_id=img_info.image_id,
                        bbox_in_image=item.bbox,
                        ocr_confidence=item.confidence,
                        polygon=item.polygon,
                        style=style,
                        is_translatable=is_target,
                        skip_reason=reason if not is_target else None,
                    )
                    slide_elements.append(element)

            # Parse presenter notes
            if self.translate_notes and slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_frame = slide.notes_slide.notes_text_frame
                for p_idx, paragraph in enumerate(notes_frame.paragraphs):
                    text, runs_info, style = merge_paragraph_runs(paragraph)
                    if not text.strip():
                        continue
                    element_id = f"slide_{slide_idx+1}_notes_p_{p_idx}"
                    element = TextElement(
                        element_id=element_id,
                        slide_index=slide_idx,
                        source_type="notes",
                        source_text=text,
                        runs=runs_info,
                        style=style,
                        paragraph_index=p_idx,
                        is_translatable=is_translatable_text(text),
                    )
                    slide_elements.append(element)

            slide_model = SlideModel(
                slide_index=slide_idx,
                slide_id=slide.slide_id,
                title=slide_title,
                elements=slide_elements,
            )
            slides.append(slide_model)

        return IntermediateRepresentation(
            title=presentation_title,
            slides=slides,
        )

    def _get_ocr_for_image(self, img_info) -> List[OCRResultItem]:
        cached = self.cache_manager.get_ocr(img_info.image_hash, self.ocr_lang)
        if cached is not None:
            return [OCRResultItem(**item) for item in cached]

        ocr_items = self.ocr_detector.detect(img_info.image_bytes)
        self.cache_manager.set_ocr(
            img_info.image_hash,
            self.ocr_lang,
            [item.to_dict() for item in ocr_items],
        )
        return ocr_items

    def _parse_shape(self, shape, slide_idx: int, slide_elements: List[TextElement]) -> None:
        shape_id = shape.shape_id if hasattr(shape, "shape_id") else 0

        # Group shapes
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub_shape in shape.shapes:
                self._parse_shape(sub_shape, slide_idx, slide_elements)
            return

        # Table
        if shape.has_table:
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                for col_idx, cell in enumerate(row.cells):
                    if cell.text_frame:
                        for p_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                            text, runs_info, style = merge_paragraph_runs(paragraph)
                            if not text.strip():
                                continue
                            element_id = f"slide_{slide_idx+1}_table_{shape_id}_r{row_idx}_c{col_idx}_p{p_idx}"
                            bbox = [
                                float(shape.left),
                                float(shape.top),
                                float(shape.width),
                                float(shape.height),
                            ]
                            element = TextElement(
                                element_id=element_id,
                                slide_index=slide_idx,
                                source_type="table_cell",
                                source_text=text,
                                bbox=bbox,
                                style=style,
                                shape_id=shape_id,
                                paragraph_index=p_idx,
                                cell_address=[row_idx, col_idx],
                                runs=runs_info,
                                bullet_level=paragraph.level,
                                is_translatable=is_translatable_text(text),
                            )
                            slide_elements.append(element)
            return

        # Text Frame / AutoShape text
        if shape.has_text_frame and shape.text_frame:
            for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                text, runs_info, style = merge_paragraph_runs(paragraph)
                if not text.strip():
                    continue
                element_id = f"slide_{slide_idx+1}_shape_{shape_id}_p_{p_idx}"
                bbox = [
                    float(shape.left) if hasattr(shape, "left") else 0.0,
                    float(shape.top) if hasattr(shape, "top") else 0.0,
                    float(shape.width) if hasattr(shape, "width") else 0.0,
                    float(shape.height) if hasattr(shape, "height") else 0.0,
                ]
                element = TextElement(
                    element_id=element_id,
                    slide_index=slide_idx,
                    source_type="pptx_text",
                    source_text=text,
                    bbox=bbox,
                    style=style,
                    shape_id=shape_id,
                    paragraph_index=p_idx,
                    runs=runs_info,
                    bullet_level=paragraph.level,
                    is_translatable=is_translatable_text(text),
                )
                slide_elements.append(element)
