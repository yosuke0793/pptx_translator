import logging
from typing import Any, Dict

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_translator.models.element import TextElement
from pptx_translator.models.slide import IntermediateRepresentation

logger = logging.getLogger(__name__)


class TextReplacer:
    """Replaces original text in PPTX with translations in-place."""

    def __init__(self, pptx_path: str):
        self.pptx_path = pptx_path
        self.prs = Presentation(pptx_path)

    def replace_texts(self, ir: IntermediateRepresentation) -> Any:
        # Build lookup map of element_id -> TextElement
        elem_map: Dict[str, TextElement] = {}
        for slide in ir.slides:
            for elem in slide.elements:
                elem_map[elem.element_id] = elem

        for slide_idx, slide in enumerate(self.prs.slides):
            # Process shapes
            for shape in slide.shapes:
                self._process_shape(shape, slide_idx, elem_map)

            # Process notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_frame = slide.notes_slide.notes_text_frame
                for p_idx, paragraph in enumerate(notes_frame.paragraphs):
                    elem_id = f"slide_{slide_idx+1}_notes_p_{p_idx}"
                    if elem_id in elem_map and elem_map[elem_id].translation:
                        self._apply_translation_to_paragraph(paragraph, elem_map[elem_id])

        return self.prs

    def _process_shape(self, shape, slide_idx: int, elem_map: Dict[str, TextElement]) -> None:
        shape_id = shape.shape_id if hasattr(shape, "shape_id") else 0

        # Group shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub_shape in shape.shapes:
                self._process_shape(sub_shape, slide_idx, elem_map)
            return

        # Table
        if shape.has_table:
            table = shape.table
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    if cell.text_frame:
                        for p_idx, paragraph in enumerate(cell.text_frame.paragraphs):
                            elem_id = f"slide_{slide_idx+1}_table_{shape_id}_r{r_idx}_c{c_idx}_p{p_idx}"
                            if elem_id in elem_map and elem_map[elem_id].translation:
                                self._apply_translation_to_paragraph(paragraph, elem_map[elem_id])
            return

        # Regular shape text frame
        if shape.has_text_frame and shape.text_frame:
            for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                elem_id = f"slide_{slide_idx+1}_shape_{shape_id}_p_{p_idx}"
                if elem_id in elem_map and elem_map[elem_id].translation:
                    self._apply_translation_to_paragraph(paragraph, elem_map[elem_id])

    def _apply_translation_to_paragraph(self, paragraph, elem: TextElement) -> None:
        new_text = elem.translation or elem.source_text
        if not paragraph.runs:
            paragraph.text = new_text
            return

        # Preserve first run's font formatting if possible, clear remaining runs
        first_run = paragraph.runs[0]
        font_name = first_run.font.name
        font_size = first_run.font.size
        bold = first_run.font.bold
        italic = first_run.font.italic

        # Clear text from all runs
        for run in paragraph.runs:
            run.text = ""

        # Set translation on first run
        first_run.text = new_text
        if font_name:
            first_run.font.name = font_name
        if font_size:
            first_run.font.size = font_size
        if bold is not None:
            first_run.font.bold = bold
        if italic is not None:
            first_run.font.italic = italic
