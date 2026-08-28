import logging
from typing import Any

from pptx.dml.color import RGBColor
from pptx.util import Pt

from pptx_translator.models.slide import IntermediateRepresentation

logger = logging.getLogger(__name__)


def parse_hex_color(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    if len(hex_str) == 6:
        r = int(hex_str[0:2], 16)
        g = int(hex_str[2:4], 16)
        b = int(hex_str[4:6], 16)
        return RGBColor(r, g, b)
    return RGBColor(0, 0, 0)


class OverlayTextWriter:
    """Adds overlay text boxes on slides corresponding to translated image texts."""

    def add_overlays(self, prs: Any, ir: IntermediateRepresentation) -> Any:
        for slide_model in ir.slides:
            slide_idx = slide_model.slide_index
            if slide_idx >= len(prs.slides):
                continue
            slide = prs.slides[slide_idx]

            # Collect existing shape names on slide to avoid duplication on rerun
            existing_names = {shape.name for shape in slide.shapes if hasattr(shape, "name")}

            image_text_elements = [
                elem
                for elem in slide_model.elements
                if elem.source_type == "image_text" and elem.translation and elem.bbox
            ]

            for elem in image_text_elements:
                if elem.element_id in existing_names:
                    logger.info(f"Overlay textbox {elem.element_id} already exists on slide {slide_idx+1}. Skipping.")
                    continue

                if not elem.bbox or len(elem.bbox) < 4:
                    continue
                xp, yp, wp, hp = elem.bbox
                tx_box = slide.shapes.add_textbox(int(xp), int(yp), int(wp), int(hp))
                tx_box.name = elem.element_id  # Set unique identifier name

                tf = tx_box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = elem.translation or ""

                # Styling
                if elem.style.color:
                    try:
                        p.font.color.rgb = parse_hex_color(elem.style.color)
                    except Exception:
                        pass
                if elem.style.font_size:
                    p.font.size = Pt(elem.style.font_size)
                else:
                    p.font.size = Pt(12)

                logger.info(f"Added overlay textbox {elem.element_id} on slide {slide_idx+1}")

        return prs
