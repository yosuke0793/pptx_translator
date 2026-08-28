import numpy as np
from pptx import Presentation

from pptx_translator.image_processing.background_cleaner import BackgroundCleaner
from pptx_translator.image_processing.coordinate_transform import image_to_slide_coords
from pptx_translator.image_processing.mask_generator import MaskGenerator
from pptx_translator.models.element import TextElement
from pptx_translator.models.slide import IntermediateRepresentation, SlideModel
from pptx_translator.writer.overlay_text_writer import OverlayTextWriter


def test_coordinate_transform():
    # Image 100x100 pixels mapped to slide 1000x1000 EMU
    coords = image_to_slide_coords(
        image_bbox=[10, 20, 30, 40],
        pixel_size=(100, 100),
        slide_bbox=[100, 200, 1000, 1000],
    )
    # x_p = 100 + (10/100)*1000 = 200
    # y_p = 200 + (20/100)*1000 = 400
    # w_p = (30/100)*1000 = 300
    # h_p = (40/100)*1000 = 400
    assert coords == [200.0, 400.0, 300.0, 400.0]


def test_coordinate_transform_with_crop():
    # Crop 10% on left and right: effective width is 80px
    coords = image_to_slide_coords(
        image_bbox=[20, 0, 40, 50],
        pixel_size=(100, 100),
        slide_bbox=[0, 0, 800, 800],
        crop=[0.1, 0.0, 0.1, 0.0],
    )
    # rel_x = 20 - 10 = 10. eff_pw = 80. xp = (10/80)*800 = 100
    assert coords[0] == 100.0


def test_mask_generator_and_cleaner():
    img_np = np.ones((100, 100, 3), dtype=np.uint8) * 255
    # Draw black text rectangle
    img_np[20:40, 20:40] = 0

    mask_gen = MaskGenerator(dilation_pixels=2)
    mask = mask_gen.create_mask((100, 100, 3), [[20, 20, 20, 20]])
    assert mask[25, 25] == 255

    cleaner = BackgroundCleaner(method="solid")
    cleaned = cleaner.clean_image(img_np, mask)
    # Text region should be cleaned (filled with white)
    assert cleaned[25, 25, 0] > 200


def test_overlay_text_writer_deduplication():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    ir = IntermediateRepresentation(
        slides=[
            SlideModel(
                slide_index=0,
                slide_id=1,
                elements=[
                    TextElement(
                        element_id="slide_1_image_1_text_1",
                        slide_index=0,
                        source_type="image_text",
                        source_text="テスト",
                        translation="Test Overlay",
                        bbox=[100, 100, 200, 50],
                    )
                ],
            )
        ]
    )

    writer = OverlayTextWriter()
    prs = writer.add_overlays(prs, ir)
    assert len(slide.shapes) == 1
    assert slide.shapes[0].name == "slide_1_image_1_text_1"

    # Second pass - should not duplicate overlay
    prs = writer.add_overlays(prs, ir)
    assert len(slide.shapes) == 1
