import os
import tempfile

from pptx import Presentation
from pptx.util import Inches

from pptx_translator.parser.pptx_parser import PPTXParser
from pptx_translator.parser.text_run_merger import has_japanese, is_translatable_text
from pptx_translator.translation.dummy_provider import DummyProvider
from pptx_translator.writer.text_replacer import TextReplacer


def create_sample_pptx(path: str):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "日本語タイトルテスト"
    subtitle.text = "これはサブタイトルの日本語テキストです。http://example.com"

    # Add a table slide
    blank_layout = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(blank_layout)
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "項目名"
    table.cell(0, 1).text = "解析結果"
    table.cell(1, 0).text = "English Only"
    table.cell(1, 1).text = "応力分布"

    prs.save(path)


def test_has_japanese_and_translatable():
    assert has_japanese("こんにちは")
    assert has_japanese("テスト 123")
    assert not has_japanese("Hello World")

    assert is_translatable_text("日本語のテキスト")
    assert not is_translatable_text("https://example.com")
    assert not is_translatable_text("user@example.com")
    assert not is_translatable_text("12345")
    assert not is_translatable_text("English text only")


def test_pptx_parser_and_replacer():
    with tempfile.TemporaryDirectory() as tmpdir:
        input_pptx = os.path.join(tmpdir, "test_input.pptx")
        output_pptx = os.path.join(tmpdir, "test_output.pptx")

        create_sample_pptx(input_pptx)

        parser = PPTXParser(input_pptx)
        ir = parser.parse()

        assert len(ir.slides) == 2
        elements = [elem for slide in ir.slides for elem in slide.elements]
        assert len(elements) >= 5

        # Perform dummy translation
        provider = DummyProvider(prefix="[EN] ")
        ir = provider.translate_all(ir)

        # Replace text
        replacer = TextReplacer(input_pptx)
        out_prs = replacer.replace_texts(ir)
        out_prs.save(output_pptx)

        assert os.path.exists(output_pptx)

        # Parse translated PPTX to check replacements
        parser2 = PPTXParser(output_pptx)
        ir2 = parser2.parse()
        all_texts = " ".join(elem.source_text for slide in ir2.slides for elem in slide.elements)
        assert "[EN]" in all_texts or "Overview" in all_texts or "Analysis Results" in all_texts
